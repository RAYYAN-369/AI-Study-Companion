"""
Document parsing service.

Extracts plain text from an uploaded file. Supported formats: PDF, TXT, DOCX
(kept in sync with `settings.ALLOWED_EXTENSIONS`).
"""

import os

from docx import Document
from pypdf import PdfReader


def extract_pdf(file_path: str) -> str:
    text = ""
    reader = PdfReader(file_path)

    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"

    return text


def extract_txt(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as file:
        return file.read()


def extract_docx(file_path: str) -> str:
    document = Document(file_path)
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    # Also pull text out of tables, since notes/syllabi often use them.
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    return "\n".join(paragraphs)


def extract_text(file_path: str) -> str:
    """
    Dispatch to the right extractor based on file extension.
    Raises ValueError for unsupported formats.
    """
    ext = os.path.splitext(file_path)[1].lower().lstrip(".")

    if ext == "pdf":
        return extract_pdf(file_path)
    elif ext == "txt":
        return extract_txt(file_path)
    elif ext == "docx":
        return extract_docx(file_path)
    else:
        raise ValueError(f"Unsupported file format: .{ext}")
